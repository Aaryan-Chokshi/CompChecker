import json
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock
from uuid import UUID, uuid4

import numpy as np
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader
from qdrant_client import models
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

from backend.core.errors import ApplicationError, ProviderError
from backend.features.compliance.schemas import (
    Citation,
    Evaluation,
    EvaluationBatch,
    NumberedEvaluation,
    Satisfaction,
)
from backend.features.documents.chunking import AgenticChunker, ChunkGroup, ChunkPlan
from backend.features.documents.parsing import SourceUnit, parse_document
from backend.integrations.embeddings.huggingface import HuggingFaceEmbeddings
from backend.main import create_app


class FakeLLM:
    def __init__(self):
        self.calls = []
        self.status = Satisfaction.PARTIAL
        self.confidence = 0.85
        self.invalid_citation = False

    def generate_structured(self, instructions, user_input, response_model):
        data = json.loads(user_input)
        self.calls.append((instructions, data, response_model))
        if response_model is ChunkPlan:
            return ChunkPlan(
                groups=[ChunkGroup(unit_ids=[unit["id"]], topic="Requirement") for unit in data]
            )
        if response_model is EvaluationBatch:
            return EvaluationBatch(
                evaluations=[
                    NumberedEvaluation(task_id=task["task_id"], evaluation=self._evaluation(task))
                    for task in data["tasks"]
                ]
            )
        return self._evaluation(data)

    def _evaluation(self, data):
        policy, supplier = data["policy_chunks"], data["supplier_chunks"]
        status = self.status if policy and supplier else Satisfaction.INSUFFICIENT_EVIDENCE
        extent = {
            Satisfaction.PARTIAL: 50,
            Satisfaction.SATISFIED: 100,
            Satisfaction.NOT_SATISFIED: 0,
        }.get(status)

        def cite(chunks):
            return (
                [
                    Citation(
                        chunk_id=chunks[0]["id"],
                        excerpt="invented" if self.invalid_citation else chunks[0]["text"],
                    )
                ]
                if chunks
                else []
            )

        return Evaluation(
            requirement="Minimum tensile strength",
            satisfaction=status,
            extent_percent=extent,
            confidence=self.confidence,
            explanation="Compare reported measurements with the regulatory limit.",
            policy_evidence=cite(policy),
            supplier_evidence=cite(supplier),
            clarification_questions=[],
        )


@pytest.fixture
def workflow_api(settings):
    app = create_app(settings)
    with TestClient(app) as client:
        model = Mock()
        model.get_sentence_embedding_dimension.return_value = 3
        model.encode.side_effect = lambda texts, **kwargs: np.array(
            [[1.0, 0.0, 0.0] for text in texts]
        )
        app.state.providers.embeddings._model_factory = Mock(return_value=model)
        llm = FakeLLM()
        app.state.providers.llm.generate_structured = llm.generate_structured
        client.fake_llm = llm
        yield client


def upload(
    api,
    role,
    filename="measurements.csv",
    content=b"Parameter,Value,Unit\nStrength,90,MPa\nThickness,2,mm\n",
):
    response = api.post(
        "/api/v1/documents", files={"file": (filename, content)}, data={"role": role}
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_same_pipeline_preserves_one_vector_per_row_and_survives_restart(workflow_api, settings):
    ids = []
    for role in ("policy", "subject"):
        document = upload(workflow_api, role)
        ids.append(document["id"])
        assert document["status"] == "ready"
        chunks = workflow_api.get(f"/api/v1/documents/{document['id']}/chunks").json()
        assert len(chunks) == 2
        assert "Strength" in chunks[0]["text"] and "Thickness" not in chunks[0]["text"]
        assert chunks[0]["provenance"]["row"] == 2
        assert len(chunks[0]["provenance"]["columns"]) == 3
        store = workflow_api.app.state.providers.vector_store
        hits = store.query(
            [1, 0, 0],
            query_filter=models.Filter(
                must=[
                    models.FieldCondition(
                        key="document_id", match=models.MatchValue(value=document["id"])
                    )
                ]
            ),
        )
        assert {hit.id for hit in hits} == {c["id"] for c in chunks}
        assert all(hit.payload["role"] == role for hit in hits)
    assert (
        not workflow_api.fake_llm.calls
    )  # structural rows are immutable, no semantic grouping needed
    with TestClient(create_app(settings)) as restarted:
        for document_id in ids:
            assert len(restarted.get(f"/api/v1/documents/{document_id}/chunks").json()) == 2


def test_agentic_prose_is_verbatim_and_table_boundaries_are_atomic(workflow_api):
    content = b"Material must resist heat. Supplier must provide a certificate.\n\n| Parameter | Value |\n| --- | --- |\n| Strength | 100 MPa |\n| Thickness | 2 mm |\n"
    policy = upload(workflow_api, "policy", "policy.md", content)
    subject = upload(workflow_api, "subject", "subject.md", content)
    for document in (policy, subject):
        chunks = workflow_api.get(f"/api/v1/documents/{document['id']}/chunks").json()
        assert [c["provenance"]["kind"] for c in chunks] == ["prose", "prose", "row", "row"]
        assert chunks[0]["text"] == "Material must resist heat."
        assert chunks[1]["text"] == "Supplier must provide a certificate."
        assert "Thickness" not in chunks[2]["text"]
    assert len(workflow_api.fake_llm.calls) == 2


@pytest.mark.parametrize("groups", [[], [[0], [0]], [[1], [0]], [[0]], [[0, 1, 2]]])
def test_invalid_agent_plans_fall_back_without_losing_or_inventing_text(settings, groups):
    llm = Mock()
    llm.generate_structured.return_value = ChunkPlan(
        groups=[ChunkGroup(unit_ids=ids, topic="x") for ids in groups]
    )
    chunks = AgenticChunker(llm, settings).chunk(
        [
            SourceUnit("First sentence.", {"kind": "prose"}),
            SourceUnit("Second sentence.", {"kind": "prose"}),
        ]
    )
    assert [chunk.text for chunk in chunks] == ["First sentence.", "Second sentence."]
    assert all(chunk.provenance["strategy"] == "deterministic_fallback" for chunk in chunks)


def test_invalid_agent_plan_does_not_fail_document_upload(workflow_api):
    workflow_api.app.state.providers.llm.generate_structured = lambda *args: ChunkPlan(groups=[])
    document = upload(
        workflow_api,
        "policy",
        "policy.txt",
        b"First requirement. Second requirement.",
    )
    chunks = workflow_api.get(f"/api/v1/documents/{document['id']}/chunks").json()
    assert document["status"] == "ready"
    assert [chunk["text"] for chunk in chunks] == [
        "First requirement.",
        "Second requirement.",
    ]
    assert all(chunk["provenance"]["strategy"] == "deterministic_fallback" for chunk in chunks)


def test_docx_and_pdf_tables_do_not_duplicate_rows_as_prose():
    word = Document()
    word.add_paragraph("Technical specifications")
    table = word.add_table(rows=3, cols=2)
    for row, values in zip(
        table.rows,
        [["Parameter", "Value"], ["Strength", "100 MPa"], ["Thickness", "2 mm"]],
        strict=True,
    ):
        for cell, text in zip(row.cells, values, strict=True):
            cell.text = text
    buffer = BytesIO()
    word.save(buffer)
    units = parse_document("spec.docx", buffer.getvalue())
    assert len([u for u in units if u.provenance["kind"] == "row"]) == 2
    assert "Strength" not in units[0].text
    buffer = BytesIO()
    pdf_table = Table([["Parameter", "Value"], ["Strength", "100 MPa"], ["Thickness", "2 mm"]])
    pdf_table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 1, "black")]))
    SimpleDocTemplate(buffer, pagesize=A4).build([pdf_table])
    units = parse_document("spec.pdf", buffer.getvalue())
    assert len(units) == 2
    assert all(u.provenance["kind"] == "row" for u in units)
    assert all(u.provenance["page"] == 1 for u in units)


def test_comparison_citations_confidence_sheet_and_all_download_formats(workflow_api):
    policy = upload(workflow_api, "policy")
    subject = upload(workflow_api, "subject")
    unrelated = upload(workflow_api, "policy", "unrelated.csv", b"Parameter,Value\nSecret,999\n")
    response = workflow_api.post(
        "/api/v1/checks",
        json={
            "document_id": subject["id"],
            "policy_ids": [policy["id"]],
            "specifications": ["A certificate is required."],
        },
    )
    assert response.status_code == 201, response.text
    check = response.json()
    assert check["status"] == "completed"
    result = check["result"]
    assert result["verdict"] == "non_compliant"
    assert len(result["comparison_matrix"]) == 5
    assert {r["basis"] for r in result["comparison_matrix"]} == {
        "specification",
        "policy_requirement",
    }
    assert unrelated["id"] not in json.dumps(result)
    assert result["comparison_matrix"][0]["extent_percent"] == 50
    assert result["comparison_matrix"][0]["confidence"] == 0.85
    assert result["comparison_matrix"][0]["policy_evidence"][0]["document_id"] == policy["id"]
    assert result["clarification_questions"]
    assert "Supplier answer" in result["clarification_markdown"]
    markdown = workflow_api.get(f"/api/v1/checks/{check['id']}/clarifications")
    assert markdown.text == result["clarification_markdown"]
    word = workflow_api.get(f"/api/v1/checks/{check['id']}/clarifications?format=docx")
    pdf = workflow_api.get(f"/api/v1/checks/{check['id']}/clarifications?format=pdf")
    assert word.status_code == pdf.status_code == 200
    assert "attachment" in word.headers["content-disposition"]
    assert "Supplier clarification sheet" in Document(BytesIO(word.content)).paragraphs[0].text
    assert "Supplier clarification sheet" in PdfReader(BytesIO(pdf.content)).pages[0].extract_text()
    assert (
        workflow_api.get(f"/api/v1/checks/{check['id']}/clarifications?format=exe").status_code
        == 422
    )


def test_comparison_batches_tasks_and_embeds_retrieval_queries_once(workflow_api, settings):
    settings.comparison_batch_size = 2
    policy, subject = upload(workflow_api, "policy"), upload(workflow_api, "subject")
    model = workflow_api.app.state.providers.embeddings._model_factory.return_value
    encode_calls_before = model.encode.call_count

    response = workflow_api.post(
        "/api/v1/checks", json={"document_id": subject["id"], "policy_ids": [policy["id"]]}
    )

    assert response.status_code == 201, response.text
    assert len(response.json()["result"]["comparison_matrix"]) == 4
    assert model.encode.call_count == encode_calls_before + 1
    batch_calls = [
        data
        for _, data, response_model in workflow_api.fake_llm.calls
        if response_model is EvaluationBatch
    ]
    assert len(batch_calls) == 2
    assert [len(batch["tasks"]) for batch in batch_calls] == [2, 2]
    assert all([task["task_id"] for task in batch["tasks"]] == [0, 1] for batch in batch_calls)


def test_incomplete_comparison_batch_retries_as_smaller_batches(workflow_api, settings):
    settings.comparison_batch_size = 4
    policy, subject = upload(workflow_api, "policy"), upload(workflow_api, "subject")
    original = workflow_api.fake_llm.generate_structured

    def incomplete_batch(instructions, user_input, response_model):
        if response_model is EvaluationBatch and len(json.loads(user_input)["tasks"]) > 1:
            return EvaluationBatch(evaluations=[])
        return original(instructions, user_input, response_model)

    workflow_api.app.state.providers.llm.generate_structured = incomplete_batch
    response = workflow_api.post(
        "/api/v1/checks", json={"document_id": subject["id"], "policy_ids": [policy["id"]]}
    )

    assert response.status_code == 201, response.text
    assert len(response.json()["result"]["comparison_matrix"]) == 4
    assert len(workflow_api.fake_llm.calls) == 4


def test_hallucinated_citation_persists_failed_review(workflow_api):
    policy, subject = upload(workflow_api, "policy"), upload(workflow_api, "subject")
    workflow_api.fake_llm.invalid_citation = True
    response = workflow_api.post(
        "/api/v1/checks", json={"document_id": subject["id"], "policy_ids": [policy["id"]]}
    )
    assert response.status_code == 502
    assert response.json()["error"]["code"] == "evaluation_invalid_citation"
    failed = workflow_api.get("/api/v1/checks").json()[0]
    assert failed["status"] == "failed" and failed["result"] is None
    assert workflow_api.get(f"/api/v1/checks/{failed['id']}/clarifications").status_code == 409


def test_partial_qdrant_write_is_compensated_and_no_chunks_are_committed(workflow_api, monkeypatch):
    store = workflow_api.app.state.providers.vector_store
    original = store.upsert

    def fail_after_write(records):
        original(records)
        raise ProviderError("vector_upsert_failed", "Simulated write failure")

    monkeypatch.setattr(store, "upsert", fail_after_write)
    response = workflow_api.post(
        "/api/v1/documents",
        files={"file": ("spec.csv", b"Parameter,Value\nStrength,100\n")},
        data={"role": "subject"},
    )
    assert response.status_code == 502
    document = workflow_api.get("/api/v1/documents").json()[0]
    assert document["status"] == "failed"
    assert workflow_api.get(f"/api/v1/documents/{document['id']}/chunks").json() == []
    assert store.query([1, 0, 0]) == []


@pytest.mark.parametrize(
    "filename,content,status",
    [
        ("bad.exe", b"data", 415),
        ("empty.txt", b"", 422),
        ("bad.pdf", b"not a pdf", 422),
        ("bad.txt", b"\xff", 422),
    ],
)
def test_invalid_uploads_do_not_initialize_providers(api, filename, content, status):
    response = api.post(
        "/api/v1/documents", files={"file": (filename, content)}, data={"role": "subject"}
    )
    assert response.status_code == status
    assert not api.app.state.providers.embeddings.initialized
    assert api.get("/api/v1/documents").json() == []


def test_size_guards_and_ready_roles(workflow_api, settings):
    settings.max_upload_bytes = 10
    response = workflow_api.post(
        "/api/v1/documents", files={"file": ("large.txt", b"a" * 11)}, data={"role": "subject"}
    )
    assert response.status_code == 413
    settings.max_upload_bytes = 1000
    policy = upload(workflow_api, "policy")
    response = workflow_api.post(
        "/api/v1/checks", json={"document_id": policy["id"], "policy_ids": [policy["id"]]}
    )
    assert response.status_code == 409
    assert workflow_api.get("/api/v1/checks").json() == []
    assert workflow_api.get(f"/api/v1/documents/{uuid4()}/chunks").status_code == 404
    assert UUID(policy["id"])


def test_embedding_checks_cache_then_downloads_and_reuses_model(settings):
    settings.embedding_local_files_only = False
    model = Mock()
    model.get_sentence_embedding_dimension.return_value = 3
    model.encode.return_value = np.array([[1.0, 0.0, 0.0]])
    factory = Mock(side_effect=[OSError("not cached"), model])
    embeddings = HuggingFaceEmbeddings(settings, factory)
    embeddings.embed(["text"])
    embeddings.embed(["second text"])
    assert [call.kwargs["local_files_only"] for call in factory.call_args_list] == [True, False]
    assert settings.embedding_cache_dir.is_dir()


def test_embedding_token_limit_prevents_silent_row_truncation(settings):
    model = SimpleNamespace(
        max_seq_length=2, tokenizer=SimpleNamespace(encode=lambda *args, **kwargs: [1, 2, 3])
    )
    with pytest.raises(ProviderError) as exc:
        HuggingFaceEmbeddings(settings, lambda *args, **kwargs: model).embed(["long row"])
    assert exc.value.code == "embedding_input_too_long"


def test_csv_quoted_multiline_cells_and_duplicate_headers_stay_in_single_row():
    units = parse_document("quoted.csv", b'Name,Name,Value\nA,B,"First line\nSecond line"\nC,D,0\n')
    assert len(units) == 2
    assert units[0].provenance["columns"] == [
        {"name": "Name", "value": "A"},
        {"name": "Name", "value": "B"},
        {"name": "Value", "value": "First line\nSecond line"},
    ]
    assert "Name: C" not in units[0].text
    assert "Value: 0" in units[1].text


def test_changed_model_and_lowered_chunk_limit_do_not_silently_review_partial_documents(
    workflow_api, settings
):
    policy, subject = upload(workflow_api, "policy"), upload(workflow_api, "subject")
    body = {"document_id": subject["id"], "policy_ids": [policy["id"]]}
    original = settings.embedding_model_id
    settings.embedding_model_id = "another-model"
    response = workflow_api.post("/api/v1/checks", json=body)
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "embedding_model_changed"
    settings.embedding_model_id = original
    settings.max_document_chunks = 1
    response = workflow_api.post("/api/v1/checks", json=body)
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "too_many_chunks"
    assert workflow_api.get("/api/v1/checks").json() == []


@pytest.mark.parametrize(
    "status,confidence,verdict",
    [
        (Satisfaction.SATISFIED, 0.95, "compliant"),
        (Satisfaction.SATISFIED, 0.5, "needs_review"),
        (Satisfaction.INSUFFICIENT_EVIDENCE, 0.9, "needs_review"),
    ],
)
def test_matrix_verdict_and_clarifications_reflect_uncertainty(
    workflow_api, status, confidence, verdict
):
    policy, subject = upload(workflow_api, "policy"), upload(workflow_api, "subject")
    workflow_api.fake_llm.status = status
    workflow_api.fake_llm.confidence = confidence
    response = workflow_api.post(
        "/api/v1/checks", json={"document_id": subject["id"], "policy_ids": [policy["id"]]}
    )
    assert response.status_code == 201, response.text
    result = response.json()["result"]
    assert result["verdict"] == verdict
    if verdict == "compliant":
        assert not result["clarification_questions"]
        assert "No supplier clarification questions" in result["clarification_markdown"]
    else:
        assert result["clarification_questions"]


def test_oversized_atomic_row_is_rejected_without_llm(settings):
    settings.chunk_max_chars = 100
    llm = Mock()
    with pytest.raises(ApplicationError) as exc:
        AgenticChunker(llm, settings).chunk([SourceUnit("a" * 101, {"kind": "row"})])
    assert exc.value.code == "row_too_large"
    llm.generate_structured.assert_not_called()


def test_nested_and_vertically_merged_docx_tables_fail_explicitly():
    for nested in [True, False]:
        word = Document()
        table = word.add_table(rows=2, cols=2)
        if nested:
            table.cell(0, 0).add_table(rows=2, cols=2)
        else:
            table.cell(0, 0).merge(table.cell(1, 0))
        stream = BytesIO()
        word.save(stream)
        with pytest.raises(ApplicationError) as exc:
            parse_document("merged.docx", stream.getvalue())
        assert exc.value.code == "unsupported_table_layout"


@pytest.mark.parametrize("failure", ["missing", "unknown_id"])
def test_missing_or_inconsistent_vectors_fail_before_ollama(workflow_api, monkeypatch, failure):
    policy, subject = upload(workflow_api, "policy"), upload(workflow_api, "subject")
    hits = [] if failure == "missing" else [SimpleNamespace(id=str(uuid4()), score=0.9)]
    monkeypatch.setattr(
        workflow_api.app.state.providers.vector_store, "query", lambda *a, **kw: hits
    )
    response = workflow_api.post(
        "/api/v1/checks", json={"document_id": subject["id"], "policy_ids": [policy["id"]]}
    )
    assert response.status_code == 409
    assert response.json()["error"]["code"] == (
        "vector_index_missing" if failure == "missing" else "vector_index_inconsistent"
    )
    assert not workflow_api.fake_llm.calls
    assert workflow_api.get("/api/v1/checks").json()[0]["status"] == "failed"


def test_retrieved_context_survives_omitted_llm_citations(workflow_api):
    policy, subject = upload(workflow_api, "policy"), upload(workflow_api, "subject")
    original = workflow_api.fake_llm.generate_structured

    def omit_specification_citations(instructions, user_input, response_model):
        result = original(instructions, user_input, response_model)
        tasks = json.loads(user_input)["tasks"]
        for task, item in zip(tasks, result.evaluations, strict=True):
            if task["basis"] == "specification":
                item.evaluation.satisfaction = Satisfaction.INSUFFICIENT_EVIDENCE
                item.evaluation.extent_percent = None
                item.evaluation.policy_evidence = []
                item.evaluation.supplier_evidence = []
        return result

    workflow_api.app.state.providers.llm.generate_structured = omit_specification_citations
    response = workflow_api.post(
        "/api/v1/checks", json={"document_id": subject["id"], "policy_ids": [policy["id"]]}
    )
    assert response.status_code == 201, response.text
    row = response.json()["result"]["comparison_matrix"][0]
    assert not row["policy_evidence"] and not row["supplier_evidence"]
    assert len(row["retrieved_policy_chunks"]) == 2
    assert len(row["retrieved_supplier_chunks"]) == 1
    assert {c["document_id"] for c in row["retrieved_policy_chunks"]} == {policy["id"]}
    assert row["retrieved_supplier_chunks"][0]["document_id"] == subject["id"]


def test_policy_requirement_missing_citation_retries_even_for_insufficient_evidence(workflow_api):
    policy, subject = upload(workflow_api, "policy"), upload(workflow_api, "subject")
    original = workflow_api.fake_llm.generate_structured
    feedback = []

    def omit_policy_citation(instructions, user_input, response_model):
        payload = json.loads(user_input)
        result = original(instructions, user_input, response_model)
        for task, item in zip(payload["tasks"], result.evaluations, strict=True):
            if task["basis"] != "policy_requirement":
                continue
            item.evaluation.satisfaction = Satisfaction.INSUFFICIENT_EVIDENCE
            item.evaluation.extent_percent = None
            if "validation_feedback" in payload:
                feedback.append(payload["validation_feedback"])
            else:
                item.evaluation.policy_evidence = []
        return result

    workflow_api.app.state.providers.llm.generate_structured = omit_policy_citation
    response = workflow_api.post(
        "/api/v1/checks", json={"document_id": subject["id"], "policy_ids": [policy["id"]]}
    )
    assert response.status_code == 201, response.text
    reverse = [
        r
        for r in response.json()["result"]["comparison_matrix"]
        if r["basis"] == "policy_requirement"
    ]
    assert len(reverse) == 2
    assert all(r["policy_evidence"] for r in reverse)
    assert len(feedback) == 2


def test_retrieval_audit_uses_snapshot_while_backend_holds_qdrant_lock(workflow_api, settings):
    from scripts.diagnose_retrieval import audit

    policy, subject = upload(workflow_api, "policy"), upload(workflow_api, "subject")
    model = workflow_api.app.state.providers.embeddings._model_factory.return_value
    calls = model.encode.call_count
    report = audit(settings, subject["id"], policy["id"], samples=1)
    assert report["sqlite_chunks"] == report["qdrant_points"] == 4
    assert not report["missing_point_ids"]
    assert not report["orphan_point_ids"]
    assert not report["payload_mismatches"]
    assert report["query_counts"] == {"subject: 2 hits": 2, "policy: 2 hits": 2}
    assert len(report["examples"]) == 1
    assert model.encode.call_count == calls
    assert not workflow_api.fake_llm.calls
