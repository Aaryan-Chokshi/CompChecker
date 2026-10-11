"""One sheet representation shared by Markdown, Word and PDF exports."""

from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape


def sheet_lines(check_id: str, supplier_filename: str, questions):
    lines = [
        ("title", "Supplier clarification sheet"),
        ("paragraph", f"Document: {supplier_filename}"),
        ("paragraph", f"Review: {check_id}"),
        (
            "paragraph",
            "Please answer each question and attach supporting specifications, test reports or certificates. Include units, test methods and document references where relevant.",
        ),
    ]
    if not questions:
        lines.append(
            ("paragraph", "No supplier clarification questions were identified in this review.")
        )
    for index, item in enumerate(questions, 1):
        lines.extend(
            [
                ("heading", f"Question {index}"),
                ("paragraph", item.question),
                ("paragraph", f"Comparison reference: {item.comparison_id}"),
                ("paragraph", "Supplier answer: ________________________________________"),
                ("paragraph", "Supporting evidence: ____________________________________"),
            ]
        )
    return lines


def markdown_sheet(lines):
    def safe(text):
        # Content from filenames/LLM is displayed as text, never HTML or links.
        return (
            escape(text)
            .replace("\\", "\\\\")
            .replace("[", "\\[")
            .replace("]", "\\]")
            .replace("*", "\\*")
            .replace("_", "\\_")
            .replace("#", "\\#")
            .replace("`", "\\`")
        )

    return (
        "\n\n".join(
            {"title": "# ", "heading": "## ", "paragraph": ""}[kind] + safe(text)
            for kind, text in lines
        )
        + "\n"
    )


def export_sheet(lines, format: str):
    if format == "md":
        return markdown_sheet(lines).encode("utf-8"), "text/markdown; charset=utf-8"
    stream = BytesIO()
    if format == "docx":
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor

        document = Document()
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.8)
        normal = document.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(11)
        document.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
        for kind, text in lines:
            if kind == "title":
                document.add_paragraph(text, "Title")
            elif kind == "heading":
                document.add_heading(text, level=2)
            else:
                document.add_paragraph(text)
        document.save(stream)
        return (
            stream.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    # ReportLab ships Vera, avoiding missing glyphs for common European text.
    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    font_path = Path(reportlab.__file__).parent / "fonts/Vera.ttf"
    if "ClarificationVera" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("ClarificationVera", str(font_path)))
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "Body",
        fontName="ClarificationVera",
        fontSize=10,
        leading=15,
        spaceAfter=8,
        alignment=TA_LEFT,
        textColor=colors.black,
    )
    story = []
    for kind, text in lines:
        style = (
            styles["Title"]
            if kind == "title"
            else styles["Heading2"]
            if kind == "heading"
            else body
        )
        story.append(Paragraph(escape(text).replace("\n", "<br/>"), style))
        if kind == "title":
            story.append(Spacer(1, 10))
    SimpleDocTemplate(
        stream, pagesize=A4, leftMargin=50, rightMargin=50, topMargin=50, bottomMargin=50
    ).build(story)
    return stream.getvalue(), "application/pdf"
